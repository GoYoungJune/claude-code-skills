#!/usr/bin/env python3
"""CSV → 서식 Excel 변환 스크립트."""
import csv
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl")
    sys.exit(1)


HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
NUMBER_FORMAT = "#,##0"
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def is_numeric(values: list[str]) -> bool:
    """비어있지 않은 값 중 80% 이상이 숫자이면 숫자 열로 판단."""
    non_empty = [v for v in values if v.strip()]
    if not non_empty:
        return False
    numeric = sum(1 for v in non_empty if _to_num(v) is not None)
    return numeric / len(non_empty) >= 0.8


def _to_num(v: str):
    v = v.strip().replace(",", "")
    try:
        return int(v)
    except ValueError:
        pass
    try:
        return float(v)
    except ValueError:
        return None


def convert(csv_path: str) -> str:
    src = Path(csv_path).expanduser().resolve()
    if not src.exists():
        raise FileNotFoundError(f"파일 없음: {src}")

    with open(src, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))

    if not rows:
        raise ValueError("CSV가 비어 있습니다.")

    headers = rows[0]
    data_rows = rows[1:]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"

    # ── 헤더 행 ──────────────────────────────────────────
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    # ── 숫자 열 판별 ─────────────────────────────────────
    col_values = [
        [row[i] if i < len(row) else "" for row in data_rows]
        for i in range(len(headers))
    ]
    numeric_cols = {i for i, vals in enumerate(col_values) if is_numeric(vals)}

    # ── 데이터 행 ────────────────────────────────────────
    for row_idx, row in enumerate(data_rows, 2):
        for col_idx, val in enumerate(row, 1):
            num = _to_num(val) if (col_idx - 1) in numeric_cols else None
            cell = ws.cell(row=row_idx, column=col_idx, value=num if num is not None else val)
            cell.border = BORDER
            if num is not None:
                cell.number_format = NUMBER_FORMAT
                cell.alignment = Alignment(horizontal="right")

    # ── 합계 행 ─────────────────────────────────────────
    total_row = len(data_rows) + 2
    if data_rows and numeric_cols:
        ws.cell(row=total_row, column=1, value="합계").font = TOTAL_FONT
        ws.cell(row=total_row, column=1).fill = TOTAL_FILL
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.border = BORDER
            cell.fill = TOTAL_FILL
            if (col_idx - 1) in numeric_cols:
                col_letter = get_column_letter(col_idx)
                cell.value = f"=SUM({col_letter}2:{col_letter}{total_row - 1})"
                cell.number_format = NUMBER_FORMAT
                cell.font = TOTAL_FONT
                cell.alignment = Alignment(horizontal="right")

    # ── 열 너비 자동 조정 ────────────────────────────────
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            len(str(header)),
            *[len(str(r[col_idx - 1])) if col_idx - 1 < len(r) else 0 for r in data_rows],
            0,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

    ws.row_dimensions[1].height = 20

    # ── 차트 ─────────────────────────────────────────────
    if data_rows and numeric_cols:
        first_num_col = min(numeric_cols) + 1  # 1-indexed
        chart = BarChart()
        chart.type = "col"
        chart.title = headers[first_num_col - 1]
        chart.style = 10
        chart.y_axis.title = headers[first_num_col - 1]
        chart.x_axis.title = headers[0]
        chart.width = 15
        chart.height = 10

        data_ref = Reference(ws, min_col=first_num_col, min_row=1, max_row=len(data_rows) + 1)
        cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(data_rows) + 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        chart_col = get_column_letter(len(headers) + 2)
        ws.add_chart(chart, f"{chart_col}2")

    # ── 저장 ─────────────────────────────────────────────
    out_path = src.with_suffix(".xlsx")
    wb.save(out_path)
    return str(out_path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용법: python excel_report.py <csv_path>")
        sys.exit(1)
    result = convert(sys.argv[1])
    print(f"저장 완료: {result}")
